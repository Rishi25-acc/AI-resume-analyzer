from __future__ import annotations

import io
import os
from typing import Union

from pypdf import PdfReader
import docx


ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024


def is_valid_resume_file(filename: str, file_size_bytes: int) -> tuple[bool, str]:
 
    _, ext = os.path.splitext(filename.lower())

    if ext not in ALLOWED_EXTENSIONS:
        return False, f"Unsupported file type '{ext}'. Please upload a PDF or DOCX file."

    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        return False, f"File is too large ({file_size_bytes / (1024*1024):.1f} MB). Max allowed is {MAX_FILE_SIZE_MB} MB."

    if file_size_bytes == 0:
        return False, "The uploaded file appears to be empty."

    return True, "File looks valid."




def extract_text(file_obj: Union[str, io.BytesIO, "io.BufferedReader"], filename: str | None = None) -> str:

    name = filename or getattr(file_obj, "name", None) or (file_obj if isinstance(file_obj, str) else None)
    if not name:
        raise ValueError("Could not determine filename/extension for extraction.")

    _, ext = os.path.splitext(name.lower())

    if ext == ".pdf":
        return _extract_from_pdf(file_obj)
    elif ext == ".docx":
        return _extract_from_docx(file_obj)
    else:
        raise ValueError(f"Unsupported file extension: '{ext}'. Only .pdf and .docx are supported.")


def _extract_from_pdf(file_obj: Union[str, io.BytesIO]) -> str:
    """Extract text from every page of a PDF file."""
    reader = PdfReader(file_obj)
    pages_text = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        pages_text.append(page_text)

    return "\n".join(pages_text).strip()


def _extract_from_docx(file_obj: Union[str, io.BytesIO]) -> str:
    """Extract text from every paragraph (and table cell) of a DOCX file."""
    document = docx.Document(file_obj)
    parts = []

    # Paragraphs
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Tables (resumes sometimes use tables for layout, e.g. skills/experience)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts).strip()



if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python resume_parser.py <path_to_resume.pdf_or_.docx>")
        sys.exit(1)

    path = sys.argv[1]
    ok, msg = is_valid_resume_file(path, os.path.getsize(path))
    print(f"Validation: {ok} - {msg}")

    if ok:
        text = extract_text(path)
        print(f"\nExtracted {len(text)} characters.\n")
        print(text[:1000])
